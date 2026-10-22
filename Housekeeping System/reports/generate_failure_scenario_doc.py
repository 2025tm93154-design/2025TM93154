from docx import Document
from docx.shared import Pt
from pathlib import Path

out_file = Path(r"C:\Users\sivam\Downloads\BootStrap 5\Restaurant Equipment Lending\Housekeeping System\reports\Housekeeping_Failure_Scenarios.docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

doc.add_heading("Housekeeping System Failure Scenario Testing Guide", 0)
doc.add_paragraph("This guide captures the curl-based tests used to validate failure handling, validation flows, duplicate request protection, and status transitions in the housekeeping microservices stack.")

sections = [
    ("1. Service Health Check", [
        "curl http://localhost:3201/health",
        "curl http://localhost:3202/health",
        "curl http://localhost:3203/health",
        "curl http://localhost:3204/health",
        "Expected: each endpoint returns HTTP 200 when the service is running."
    ]),
    ("2. Valid Reservation Request", [
        "curl -X POST http://localhost:3201/api/v1/reservations \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -H \"Idempotency-Key: reservation-001\" \\",
        "  -d '{\"guestName\":\"Alice\",\"roomType\":\"STANDARD\",\"checkInDate\":\"2026-11-30\",\"nights\":2}'",
        "Expected: HTTP 201 and a reservation is created successfully."
    ]),
    ("3. Duplicate Reservation Replay", [
        "curl -X POST http://localhost:3201/api/v1/reservations \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -H \"Idempotency-Key: reservation-001\" \\",
        "  -d '{\"guestName\":\"Alice\",\"roomType\":\"STANDARD\",\"checkInDate\":\"2026-11-30\",\"nights\":2}'",
        "Expected: same duplicate request is replayed instead of creating a second reservation."
    ]),
    ("4. Invalid Reservation Payload", [
        "curl -X POST http://localhost:3201/api/v1/reservations \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -H \"Idempotency-Key: reservation-invalid-001\" \\",
        "  -d '{\"guestName\":\"\",\"roomType\":\"\",\"checkInDate\":\"\",\"nights\":0}'",
        "Expected: HTTP 400 with validation errors."
    ]),
    ("5. Valid Task Creation", [
        "curl -X POST http://localhost:3203/api/v1/tasks \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -H \"Idempotency-Key: task-001\" \\",
        "  -d '{\"reservationId\":\"res-101\",\"roomType\":\"STANDARD\",\"taskType\":\"room-cleaning\",\"assignedStaff\":\"Housekeeper Team A\"}'",
        "Expected: HTTP 201 and a valid housekeeping task is created."
    ]),
    ("6. Duplicate Task Replay", [
        "curl -X POST http://localhost:3203/api/v1/tasks \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -H \"Idempotency-Key: task-001\" \\",
        "  -d '{\"reservationId\":\"res-101\",\"roomType\":\"STANDARD\",\"taskType\":\"room-cleaning\",\"assignedStaff\":\"Housekeeper Team A\"}'",
        "Expected: HTTP 200 with replayed response and no duplicate task record."
    ]),
    ("7. Invalid Task Payload", [
        "curl -X POST http://localhost:3203/api/v1/tasks \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -H \"Idempotency-Key: task-invalid-001\" \\",
        "  -d '{\"reservationId\":\"\",\"roomType\":\"\",\"taskType\":\"\"}'",
        "Expected: HTTP 400 due to missing required values."
    ]),
    ("8. Update Task Status", [
        "curl -X PATCH http://localhost:3203/api/v1/tasks/<task-id>/status \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -d '{\"status\":\"IN_PROGRESS\"}'",
        "Expected: HTTP 200 and the task status changes to IN_PROGRESS."
    ]),
    ("9. Invalid Status Value", [
        "curl -X PATCH http://localhost:3203/api/v1/tasks/<task-id>/status \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -d '{\"status\":\"DONE\"}'",
        "Expected: HTTP 400 because DONE is not an allowed task status."
    ]),
    ("10. Fetch Updated Task", [
        "curl http://localhost:3203/api/v1/tasks/<task-id>",
        "Expected: the latest task record shows the updated status value."
    ]),
    ("11. Valid Notification Creation", [
        "curl -X POST http://localhost:3204/api/v1/notifications \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -d '{\"taskId\":\"<task-id>\",\"message\":\"Housekeeping update sent\"}'",
        "Expected: HTTP 201 when taskId and message are valid."
    ]),
    ("12. Invalid Notification Payload", [
        "curl -X POST http://localhost:3204/api/v1/notifications \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -d '{\"taskId\":\"\",\"message\":\"\"}'",
        "Expected: HTTP 400 due to validation of required fields."
    ]),
    ("13. Room Hold Creation", [
        "curl -X POST http://localhost:3202/api/v1/room-holds \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -d '{\"roomType\":\"STANDARD\",\"reservationId\":\"hold-001\",\"holdSeconds\":30}'",
        "Expected: room hold created successfully unless the room is already full."
    ]),
    ("14. Capacity Failure Scenario", [
        "curl -X POST http://localhost:3202/api/v1/room-holds \\",
        "  -H \"Content-Type: application/json\" \\",
        "  -d '{\"roomType\":\"STANDARD\",\"reservationId\":\"hold-002\",\"holdSeconds\":30}'",
        "Expected: HTTP 409 or conflict error when capacity is exhausted."
    ]),
    ("15. Expected Result Summary", [
        "Successful call -> HTTP 200 or 201",
        "Validation failure -> HTTP 400",
        "Duplicate request -> same idempotent replay response without duplication",
        "Capacity issue -> HTTP 409 or service-level conflict",
        "Missing task or invalid reference -> HTTP 404",
        "These outcomes confirm that the system handles failure scenarios in a controllable and traceable manner."
    ])
]

for title, lines in sections:
    doc.add_heading(title, level=1)
    for line in lines:
        doc.add_paragraph(line)

p = doc.add_paragraph()
r = p.add_run("Prepared for assignment documentation and testing evidence.")
r.bold = True

doc.save(out_file)
print(out_file)
